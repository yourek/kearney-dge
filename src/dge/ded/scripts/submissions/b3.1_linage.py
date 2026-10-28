"""Map the Purview lineage extract onto the catalogue IDs and fill the B.3.1 submission.

`Existing_Relations_full.xlsx` is a Purview lineage extract: one row per
column-to-column relation, naming each side by `<side>_db`, `<side>_table_name`
and `<side>_column_name`. Each side is mapped to the catalogue of its layer
(`CATALOGUES`): the table name, lowercased, to the catalogue's `Data Asset ID`
(sheet `B1.2 Catalogue Table`), then that asset and the column name to the
`Attribute ID` (sheet `B1.3 Catalogue Attribute`).

The table's `Data Owner` and `Data Steward` are taken for both sides and each
merged into one `Data Lineage Owner` / `Data Lineage Steward` - distinct
names, source side first.

`silver_modeled` has no catalogue. A minimal one is generated from the
lineage itself - `DED_Catalogue_Silver_Modeled.xlsx` beside the inputs, in
the same sheet layout as the others - numbering tables from
`DED_Data_Table_400` and attributes from `DED_Attribute_7001`, clear of the
ranges the real catalogues use. Once written it is reused, so the IDs stay
stable across runs; tables or columns that appear later are appended with
the next free numbers.

Its `Data Owner` and `Data Steward` columns are for filling in by hand. Where
they are blank, both are assumed from the table's neighbours in the lineage
(`infer_modeled_stewardship`): the LH_silver tables feeding it, else the
silver_modeled tables feeding it, else the gold tables it feeds, else the
silver_modeled tables it feeds - the distinct names across all of them.
Every assumption, and what it rests on, is listed on the mapping's
`Silver Modeled` sheet.

Two workbooks land in `<input dir>/output/`:

- `DED_<YYYYMMDD>_Lineage_B.3.1_Mapping.xlsx`, the working document: the
  extract with the mapped columns appended, the `Silver Modeled` catalogue as
  used, an `Unmapped` sheet listing each table or column that found no ID,
  and a `Summary` sheet with the coverage per layer;
- `DED_<YYYYMMDD>_Lineage_B.3.1.xlsx`, the submission, filled from
  `DED_*_Lineage_B.3.1_template.xlsx`: one `B.3.1 Data Linage` row per
  relationship (the extract's `relationship_guid` is one table-to-table
  lineage) and its `B.3.1 Asset Assignment` - the asset followed by each of
  its attributes the relationship carries, per side.

`Created Date` is fixed (`CREATED`); `Last Update Date` is the day before the
run date. The extract, the catalogues and the template are only ever read.

    uv run python src/dge/ded/scripts/submissions/b3.1_linage.py data/ded/linage_mapper
"""

from __future__ import annotations

import argparse
import re
import sys
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta
from pathlib import Path

import openpyxl
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter

from dge.ded.annotate import _strip_defined_names, _strip_legacy_parts

LINEAGE_FILE = "Existing_Relations_full.xlsx"
TEMPLATE_GLOB = "DED_*_Lineage_B.3.1_template.xlsx"
MAPPING_NAME = "DED_{date:%Y%m%d}_Lineage_B.3.1_Mapping.xlsx"
SUBMISSION_NAME = "DED_{date:%Y%m%d}_Lineage_B.3.1.xlsx"
TABLE_SHEET = "B1.2 Catalogue Table"
ATTRIBUTE_SHEET = "B1.3 Catalogue Attribute"

#: The lineage's database name, lowercased, to the catalogue file holding it.
#: Also the order lineages are listed in, upstream first.
CATALOGUES = {
    "oracle": "DED_Catalogue_Source.xlsx",
    "lh_bronze": "DED_Catalogue_Bronze.xlsx",
    "lh_silver": "DED_Catalogue_Silver.xlsx",
    "silver_modeled": "DED_Catalogue_Silver_Modeled.xlsx",
    "gold": "DED_Catalogue_Gold.xlsx",
}

#: The generated silver_modeled catalogue and where its numbering starts.
MODELED_DB = "silver_modeled"
MODELED_TABLE_START = 400
MODELED_ATTRIBUTE_START = 7001

#: Catalogue headers read here.
ASSET_ID = "Data Asset ID (English)"
ASSET_NAME = "Data Asset Name (English)"
OWNER = "Data Owner"
STEWARD = "Data Steward"
ATTRIBUTE_ID = "Attribute ID (English)"
ATTRIBUTE_NAME = "Attribute Name"
ATTRIBUTE_ASSET = "Data Asset ID"
ATTRIBUTE_TABLE = "Data Table Name"

SIDES = ("source", "target")
LINEAGE_OWNER = "Data Lineage Owner"
LINEAGE_STEWARD = "Data Lineage Steward"

#: The submission's sheets and the values it fixes.
LINEAGE_SHEET = "B.3.1 Data Linage"
ASSIGNMENT_SHEET = "B.3.1 Asset Assignment"
CREATED = date(2026, 10, 16)
DATE_FORMAT = "mm/dd/yyyy"
ARABIC_PLACEHOLDER = "-"
LINEAGE_NAME = "{target_db}_{source_db}_{source_table}_to_{target_table}"
LINEAGE_DESCRIPTION = ("Direct Linage Dataset to Dataset between "
                       "{source_db}.{source_table} and {target_db}.{target_table}")


def added_columns(side: str) -> dict[str, str]:
    """The columns appended for one side of the relation, by what they hold."""
    label = side.capitalize()
    return {"asset": f"{label} Data Asset ID", "attribute": f"{label} Attribute ID",
            "owner": f"{label} Data Owner", "steward": f"{label} Data Steward"}


def collapse(header: object) -> str:
    """A header with its whitespace collapsed, as the catalogues type it unevenly."""
    return " ".join(str(header).split()) if header is not None else ""


def text(value: object) -> str:
    return str(value).strip() if value is not None else ""


def key(value: object) -> str:
    return text(value).lower()


def read_sheet(ws, marker: str) -> list[dict[str, object]]:
    """A sheet's rows as dicts, its header being the first row holding `marker`.

    The catalogues carry a mandatory/optional row above the header; the
    generated one does not.
    """
    rows = ws.iter_rows(values_only=True)
    for row in rows:
        headers = [collapse(h) for h in row]
        if marker in headers:
            break
    else:
        raise ValueError(f"sheet {ws.title!r} has no {marker!r} header")
    # The first of a repeated header wins: Gold repeats some at the far right.
    index = {}
    for i, h in enumerate(headers):
        index.setdefault(h, i)
    return [{h: row[i] if i < len(row) else None for h, i in index.items() if h}
            for row in rows if any(v is not None for v in row)]


def table_names(value: object) -> list[str]:
    """The names a catalogue table answers to, lowercased.

    Bronze names a few tables by both names, `Investor_Dtl(B3CONTACT_ATTRIBUTE)`,
    and the lineage may use either.
    """
    name = key(value)
    found = re.fullmatch(r"(.+?)\s*\((.+)\)", name)
    return [name, *(n.strip() for n in found.groups())] if found else [name]


@dataclass
class Catalogue:
    name: str
    assets: dict[str, str] = field(default_factory=dict)  #: table name -> asset id
    owners: dict[str, str] = field(default_factory=dict)  #: asset id -> owner
    stewards: dict[str, str] = field(default_factory=dict)  #: asset id -> steward
    attributes: dict[tuple[str, str], str] = field(default_factory=dict)  #: (asset, column) -> id


def load_catalogue(path: Path) -> Catalogue:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        cat = Catalogue(path.name)
        for row in read_sheet(wb[TABLE_SHEET], ASSET_ID):
            asset = text(row.get(ASSET_ID))
            if not asset:
                continue
            for name in table_names(row.get(ASSET_NAME)):
                cat.assets.setdefault(name, asset)
            cat.owners[asset] = text(row.get(OWNER))
            cat.stewards[asset] = text(row.get(STEWARD))
        for row in read_sheet(wb[ATTRIBUTE_SHEET], ATTRIBUTE_ID):
            attribute, asset = text(row.get(ATTRIBUTE_ID)), text(row.get(ATTRIBUTE_ASSET))
            if attribute and asset:
                cat.attributes.setdefault((asset, key(row.get(ATTRIBUTE_NAME))), attribute)
        return cat
    finally:
        wb.close()


def load_lineage(path: Path) -> tuple[list[str], list[dict[str, object]]]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        rows = wb.worksheets[0].iter_rows(values_only=True)
        headers = [collapse(h) for h in next(rows)]
        return headers, [dict(zip(headers, r)) for r in rows if any(v is not None for v in r)]
    finally:
        wb.close()


def id_number(value: str) -> int:
    found = re.search(r"\d+$", value)
    return int(found.group()) if found else 0


def build_modeled_catalogue(path: Path, lineage: list[dict[str, object]]) -> tuple[int, int]:
    """Write the silver_modeled catalogue, keeping any IDs and names it already holds.

    Returns how many tables and attributes were added.
    """
    tables: list[list[object]] = []  # [asset id, name, owner, steward]
    attributes: list[list[object]] = []  # [attribute id, name, asset id, table name]
    current_layout = False
    if path.exists():
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        rows = read_sheet(wb[TABLE_SHEET], ASSET_ID)
        current_layout = not rows or OWNER in rows[0]
        tables = [[r.get(ASSET_ID), r.get(ASSET_NAME), r.get(OWNER), r.get(STEWARD)] for r in rows]
        attributes = [[r.get(ATTRIBUTE_ID), r.get(ATTRIBUTE_NAME), r.get(ATTRIBUTE_ASSET),
                       r.get(ATTRIBUTE_TABLE)] for r in read_sheet(wb[ATTRIBUTE_SHEET], ATTRIBUTE_ID)]
        wb.close()

    # Every (table, column) the lineage names on a silver_modeled side.
    wanted: dict[str, dict[str, str]] = {}
    for row in lineage:
        for side in SIDES:
            if key(row.get(f"{side}_db")) == MODELED_DB:
                table, column = text(row.get(f"{side}_table_name")), text(row.get(f"{side}_column_name"))
                if table:
                    columns = wanted.setdefault(key(table), {"": table})
                    if column:
                        columns.setdefault(key(column), column)

    assets = {key(t[1]): text(t[0]) for t in tables}
    known = {(text(a[2]), key(a[1])) for a in attributes}
    next_table = max([id_number(text(t[0])) + 1 for t in tables] + [MODELED_TABLE_START])
    next_attribute = max([id_number(text(a[0])) + 1 for a in attributes] + [MODELED_ATTRIBUTE_START])
    added_tables = added_attributes = 0
    for table in sorted(wanted):
        columns = wanted[table]
        if table not in assets:
            assets[table] = f"DED_Data_Table_{next_table:03d}"
            tables.append([assets[table], columns[""], None, None])
            next_table += 1
            added_tables += 1
        for column in sorted(c for c in columns if c):
            if (assets[table], column) not in known:
                attributes.append([f"DED_Attribute_{next_attribute:04d}", columns[column],
                                   assets[table], columns[""]])
                known.add((assets[table], column))
                next_attribute += 1
                added_attributes += 1

    if added_tables or added_attributes or not current_layout:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = TABLE_SHEET
        write_table(ws, [ASSET_ID, ASSET_NAME, OWNER, STEWARD], tables)
        write_table(wb.create_sheet(ATTRIBUTE_SHEET),
                    [ATTRIBUTE_ID, ATTRIBUTE_NAME, ATTRIBUTE_ASSET, ATTRIBUTE_TABLE], attributes)
        wb.save(path)
    return added_tables, added_attributes


def split_contacts(value: str) -> list[str]:
    return [s.strip() for s in re.split(r"[,;\n]", value) if s.strip()]


def merge_contacts(*values: str) -> str:
    """The distinct names across `values`, in first-seen order."""
    seen: dict[str, str] = {}
    for value in values:
        for contact in split_contacts(value):
            seen.setdefault(contact.lower(), contact)
    return ", ".join(seen.values())


def infer_modeled_stewardship(lineage: list[dict[str, object]],
                              catalogues: dict[str, Catalogue | None]) -> dict[str, str]:
    """Fill the silver_modeled catalogue's blank owners and stewards from its neighbours.

    Each table takes the distinct owners and stewards of its neighbours, by
    the first of these that gives any: the LH_silver tables feeding it, the
    silver_modeled tables feeding it, the gold tables it feeds, the
    silver_modeled tables it feeds. Rounds repeat until nothing changes, so
    an assumption can rest on another made in an earlier round - but never
    on one made in the same round, so the order tables are visited in does
    not matter.

    Returns the basis of each owner and steward, by asset id.
    """
    modeled = catalogues.get(MODELED_DB)
    if modeled is None:
        return {}
    upstream: dict[str, set[tuple[str, str]]] = defaultdict(set)
    downstream: dict[str, set[tuple[str, str]]] = defaultdict(set)
    for row in lineage:
        source = (key(row.get("source_db")), key(row.get("source_table_name")))
        target = (key(row.get("target_db")), key(row.get("target_table_name")))
        if target[0] == MODELED_DB:
            upstream[target[1]].add(source)
        if source[0] == MODELED_DB:
            downstream[source[1]].add(target)

    basis = {asset: "Silver Modeled catalogue" for asset in modeled.assets.values()
             if modeled.owners.get(asset) or modeled.stewards.get(asset)}
    strategies = ((upstream, "lh_silver", "fed by"), (upstream, MODELED_DB, "fed by"),
                  (downstream, "gold", "feeds"), (downstream, MODELED_DB, "feeds"))

    def known(db: str, table: str, settled: set[str]) -> tuple[str, str] | None:
        cat = catalogues.get(db)
        asset = cat.assets.get(table) if cat else None
        if asset is None or (db == MODELED_DB and asset not in settled):
            return None
        owner, steward = cat.owners.get(asset, ""), cat.stewards.get(asset, "")
        return (owner, steward) if owner or steward else None

    while True:
        settled = set(basis)
        found = {}
        for table, asset in modeled.assets.items():
            if asset in settled:
                continue
            for edges, db, relation in strategies:
                neighbours = sorted(t for d, t in edges.get(table, ()) if d == db
                                    and known(d, t, settled))
                if neighbours:
                    values = [known(db, t, settled) for t in neighbours]
                    found[asset] = (merge_contacts(*(v[0] for v in values)),
                                    merge_contacts(*(v[1] for v in values)),
                                    f"assumed: {relation} {db} " + ", ".join(neighbours))
                    break
        if not found:
            return basis
        for asset, (owner, steward, why) in found.items():
            modeled.owners[asset], modeled.stewards[asset], basis[asset] = owner, steward, why


def map_side(row: dict[str, object], side: str, catalogues: dict[str, Catalogue | None],
             unmapped: Counter[tuple[str, ...]]) -> dict[str, str]:
    """The asset id, attribute id, owner and steward for one side of a relation."""
    db, table, column = (row.get(f"{side}_{part}") for part in ("db", "table_name", "column_name"))
    cat = catalogues.get(key(db))
    out = {"asset": "", "attribute": "", "owner": "", "steward": ""}
    where = (side, text(db), text(table))
    if key(db) not in catalogues:
        unmapped[(*where, "", f"no catalogue for database {text(db)!r}")] += 1
    elif cat is None:
        unmapped[(*where, "", f"catalogue {CATALOGUES[key(db)]} not loaded")] += 1
    elif (asset := cat.assets.get(key(table))) is None:
        unmapped[(*where, "", f"table not in {cat.name}")] += 1
    else:
        out["asset"] = asset
        out["owner"], out["steward"] = cat.owners.get(asset, ""), cat.stewards.get(asset, "")
        attribute = cat.attributes.get((asset, key(column)))
        if attribute is None:
            unmapped[(*where, text(column), f"column not under {asset} in {cat.name}")] += 1
        else:
            out["attribute"] = attribute
    return out


def write_table(ws, headers: list[str], rows: list[list[object]]) -> None:
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row in rows:
        ws.append(row)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(rows) + 1}"
    for i, h in enumerate(headers, 1):
        ws.column_dimensions[get_column_letter(i)].width = max(12, min(45, len(h) + 4))


@dataclass
class Relation:
    """One table-to-table lineage: every column relation under one guid."""
    guid: str
    row: dict[str, object]  #: its first column relation, for the table-level fields
    mapped: list[dict[str, dict[str, str]]] = field(default_factory=list)  #: per column, per side

    def field(self, part: str) -> str:
        return text(self.row.get(part))

    def ids(self, side: str) -> str:
        """The side's asset followed by each distinct attribute the relation carries."""
        asset = next((m[side]["asset"] for m in self.mapped if m[side]["asset"]), "")
        attributes = dict.fromkeys(m[side]["attribute"] for m in self.mapped if m[side]["attribute"])
        return ", ".join([asset, *attributes]) if asset else ""

    def contacts(self, kind: str) -> str:
        return merge_contacts(*(m[s][kind] for m in self.mapped for s in SIDES))


def relations(lineage: list[dict[str, object]],
              mapped: list[dict[str, dict[str, str]]]) -> list[Relation]:
    """The extract grouped by relationship, upstream layers first."""
    grouped: dict[str, Relation] = {}
    for row, m in zip(lineage, mapped):
        guid = text(row.get("relationship_guid"))
        grouped.setdefault(guid, Relation(guid, row)).mapped.append(m)
    order = {db: i for i, db in enumerate(CATALOGUES)}
    return sorted(grouped.values(), key=lambda r: (
        order.get(key(r.field("source_db")), len(order)), order.get(key(r.field("target_db")), len(order)),
        key(r.field("source_table_name")), key(r.field("target_table_name"))))


def write_rows(ws, records: list[dict[str, object]]) -> None:
    """Replace the sheet's rows below the header with `records`, by header name."""
    headers = [collapse(c.value) for c in ws[1]]
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row)
    for r, record in enumerate(records, 2):
        for c, h in enumerate(headers, 1):
            value = record.get(h)
            cell = ws.cell(r, c, value if value != "" else None)
            if isinstance(value, date):
                cell.number_format = DATE_FORMAT


def write_submission(template: Path, groups: list[Relation], updated: date, out: Path) -> None:
    wb = openpyxl.load_workbook(template)
    # The template carries the catalogues' legacy names and external links;
    # saved as they are, Excel offers to repair the file.
    _strip_defined_names(wb)
    _strip_legacy_parts(wb)
    lineages, assignments = [], []
    for g in groups:
        names = {"source_db": g.field("source_db"), "target_db": g.field("target_db"),
                 "source_table": g.field("source_table_name"),
                 "target_table": g.field("target_table_name")}
        lineages.append({
            "Data Lineage ID (English)": g.guid,
            "Data Lineage Name (English)": LINEAGE_NAME.format(**names),
            "Data Lineage Name (Arabic)": ARABIC_PLACEHOLDER,
            "Data Lineage Description (English)": LINEAGE_DESCRIPTION.format(**names),
            "Data Lineage Description (Arabic)": ARABIC_PLACEHOLDER,
            "Data Lineage Owner": g.contacts("owner"),
            "Data Lineage Steward": g.contacts("steward"),
            "Created Date": CREATED,
            "Last Update Date": updated,
        })
        assignments.append({
            "Data Lineage ID": g.guid,
            "Source Data Asset IDs": g.ids("source"),
            "Target Data Asset IDs": g.ids("target"),
            "Processing Name": g.field("relationship_type"),
        })
    write_rows(wb[LINEAGE_SHEET], lineages)
    write_rows(wb[ASSIGNMENT_SHEET], assignments)
    wb.save(out)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("folder", type=Path,
                        help=f"folder holding {LINEAGE_FILE}, the DED_Catalogue_*.xlsx files "
                             f"and the submission template")
    parser.add_argument("--date", type=lambda s: datetime.strptime(s, "%Y-%m-%d").date(),
                        default=date.today(),
                        help="run date (YYYY-MM-DD): names the outputs; the day before is "
                             "the Last Update Date")
    args = parser.parse_args(argv)

    lineage_path = args.folder / LINEAGE_FILE
    templates = sorted(args.folder.glob(TEMPLATE_GLOB))
    if not lineage_path.exists():
        print(f"error: {lineage_path} not found", file=sys.stderr)
        return 1
    if not templates:
        print(f"error: no {TEMPLATE_GLOB} in {args.folder}", file=sys.stderr)
        return 1
    headers, lineage = load_lineage(lineage_path)
    missing = [f"{s}_{p}" for s in SIDES for p in ("db", "table_name", "column_name")]
    missing = [h for h in missing + ["relationship_guid", "relationship_type"] if h not in headers]
    if missing:
        print(f"error: {lineage_path.name} lacks columns {', '.join(missing)}", file=sys.stderr)
        return 1

    modeled = args.folder / CATALOGUES[MODELED_DB]
    added_tables, added_attributes = build_modeled_catalogue(modeled, lineage)
    print(f"{modeled.name}: {added_tables} tables, {added_attributes} attributes added")

    catalogues: dict[str, Catalogue | None] = {}
    for db, name in CATALOGUES.items():
        try:
            catalogues[db] = load_catalogue(args.folder / name)
        except (OSError, KeyError, ValueError) as exc:
            print(f"warning: {name}: {exc} - {db} sides left unmapped", file=sys.stderr)
            catalogues[db] = None
    basis = infer_modeled_stewardship(lineage, catalogues)

    unmapped: Counter[tuple[str, ...]] = Counter()
    coverage: Counter[tuple[str, str, str]] = Counter()  # (side, db, outcome)
    new_headers = ([h for s in SIDES for h in added_columns(s).values()]
                   + [LINEAGE_OWNER, LINEAGE_STEWARD])
    mapped_rows, out_rows = [], []
    for row in lineage:
        mapped = {s: map_side(row, s, catalogues, unmapped) for s in SIDES}
        for s in SIDES:
            outcome = ("mapped" if mapped[s]["attribute"]
                       else "table only" if mapped[s]["asset"] else "unmapped")
            coverage[(s, text(row.get(f"{s}_db")), outcome)] += 1
        extra = [mapped[s][k] or None for s in SIDES for k in added_columns(s)]
        contacts = [merge_contacts(*(mapped[s][k] for s in SIDES)) or None
                    for k in ("owner", "steward")]
        mapped_rows.append(mapped)
        out_rows.append([row.get(h) for h in headers] + extra + contacts)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Lineage"
    write_table(ws, headers + new_headers, out_rows)
    cat = catalogues.get(MODELED_DB)
    if cat is not None:
        write_table(wb.create_sheet("Silver Modeled"),
                    [ASSET_ID, ASSET_NAME, OWNER, STEWARD, "Basis"],
                    [[asset, name, cat.owners.get(asset) or None, cat.stewards.get(asset) or None,
                      basis.get(asset, "no neighbour with an owner or steward")]
                     for name, asset in sorted(cat.assets.items(), key=lambda a: a[1])])
    write_table(wb.create_sheet("Unmapped"),
                ["Side", "Database", "Table", "Column", "Reason", "Relations"],
                [[*k, n] for k, n in sorted(unmapped.items())])
    dbs = sorted({(s, db) for s, db, _ in coverage})
    write_table(wb.create_sheet("Summary"),
                ["Side", "Database", "Relations", "Mapped", "Table only", "Unmapped"],
                [[s, db, sum(coverage[(s, db, o)] for o in ("mapped", "table only", "unmapped")),
                  coverage[(s, db, "mapped")], coverage[(s, db, "table only")],
                  coverage[(s, db, "unmapped")]] for s, db in dbs])

    out_dir = args.folder / "output"
    out_dir.mkdir(exist_ok=True)
    mapping_out = out_dir / MAPPING_NAME.format(date=args.date)
    submission_out = out_dir / SUBMISSION_NAME.format(date=args.date)
    groups = relations(lineage, mapped_rows)
    try:
        wb.save(mapping_out)
        write_submission(templates[-1], groups, args.date - timedelta(days=1), submission_out)
    except PermissionError as exc:
        print(f"error: {exc.filename} is locked - close it in Excel and rerun", file=sys.stderr)
        return 1

    print(f"{len(lineage)} relations -> {mapping_out}")
    print(f"  {'side':<7}{'database':<16}{'mapped':>8}{'table only':>12}{'unmapped':>10}")
    for s, db in dbs:
        print(f"  {s:<7}{db:<16}{coverage[(s, db, 'mapped')]:>8}"
              f"{coverage[(s, db, 'table only')]:>12}{coverage[(s, db, 'unmapped')]:>10}")
    if cat is not None:
        assets = set(cat.assets.values())
        print(f"  silver_modeled owner/steward: {len(assets & basis.keys())} of "
              f"{len(assets)} tables filled - see the Silver Modeled sheet")
    print(f"  {sum(1 for r in out_rows if not r[-1])} relations without a lineage steward")
    if unmapped:
        print(f"  {len(unmapped)} distinct unmapped tables/columns - see the Unmapped sheet")
    print(f"{len(groups)} lineages from {templates[-1].name} -> {submission_out}")
    no_ids = [g for g in groups if not g.ids("source") or not g.ids("target")]
    if no_ids:
        print(f"  {len(no_ids)} lineages missing a side's asset IDs")
    return 0


if __name__ == "__main__":
    sys.exit(main())
