"""Reading catalogue sheets and the lookup workbooks they resolve against."""

from __future__ import annotations

import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Iterator

import openpyxl


def _clean(value: Any) -> Any:
    """Normalise a cell: strip whitespace, collapse blanks to None."""
    if isinstance(value, str):
        # Excel exports carry non-breaking spaces in pasted headers.
        text = value.replace("\xa0", " ").strip()
        return text or None
    return value


def clean_header(value: Any) -> str:
    """A header as the specs name it.

    Runs of whitespace are collapsed, so a column typed as
    'CDE Flag  (Y/N)' in one deliverable and 'CDE Flag (Y/N)' in another is
    the same column to every check.
    """
    text = _clean(value)
    return "" if text is None else re.sub(r"\s+", " ", str(text))


@dataclass
class Row:
    """One data row, addressable by column name."""

    number: int  # 1-indexed Excel row, so findings point at the real cell
    values: dict[str, Any]
    formats: dict[str, str] = field(default_factory=dict)

    def get(self, column: str) -> Any:
        return self.values.get(column)

    def number_format(self, column: str) -> str:
        """The cell's Excel display format, e.g. 'mm/dd/yyyy'."""
        return self.formats.get(column, "")

    def text(self, column: str) -> str:
        value = self.values.get(column)
        return "" if value is None else str(value).strip()


@dataclass
class Sheet:
    """A loaded worksheet: its headers and its data rows."""

    name: str
    headers: list[str]
    rows: list[Row]

    def has(self, column: str) -> bool:
        return column in self.headers

    def __iter__(self) -> Iterator[Row]:
        return iter(self.rows)

    def __len__(self) -> int:
        return len(self.rows)


def load_sheet(path: Path, sheet_name: str, header_row: int, first_data_row: int) -> Sheet:
    """Read one worksheet. Rows that are entirely blank are dropped."""
    book = openpyxl.load_workbook(path, data_only=True, read_only=True)
    if sheet_name not in book.sheetnames:
        raise KeyError(f"sheet {sheet_name!r} not found in {path.name}")
    worksheet = book[sheet_name]

    headers: list[str] = []
    rows: list[Row] = []
    # Cells rather than bare values, so date columns keep their display format.
    for number, raw in enumerate(worksheet.iter_rows(), start=1):
        if number == header_row:
            headers = [clean_header(cell.value) for cell in raw]
        elif number >= first_data_row:
            values: dict[str, Any] = {}
            formats: dict[str, str] = {}
            for i, cell in enumerate(raw):
                if i >= len(headers) or not headers[i]:
                    continue
                values[headers[i]] = _clean(cell.value)
                formats[headers[i]] = cell.number_format or ""
            if any(v is not None for v in values.values()):
                rows.append(Row(number, values, formats))

    book.close()
    if not headers:
        raise ValueError(f"no header found on row {header_row} of {sheet_name!r}")
    return Sheet(sheet_name, headers, rows)


def load_reference(path: Path, config: dict[str, Any]) -> list[dict[str, Any]]:
    """Read a lookup workbook (Domains, glossary) as a list of row dicts."""
    sheet = load_sheet(
        path,
        config["sheet"],
        config["header_row"],
        config["header_row"] + 1,
    )
    return [row.values for row in sheet]
